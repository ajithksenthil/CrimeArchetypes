from docx.api import Document
import xlsxwriter

# Paths for input DOCX file and output Excel file
docx_file_path = 'data/Albright_Charles.docx'
excel_file_path = 'data/Albright_Charles.xlsx'

def docx_to_xlsx(docx_file_path, excel_file_path):
    # Load the DOCX document
    document = Document(docx_file_path)
    tables = document.tables

    # Create an Excel workbook and add a worksheet
    workbook = xlsxwriter.Workbook(excel_file_path)
    worksheet = workbook.add_worksheet("Combined Data")

    # Index for tracking rows in the worksheet
    index_row = 0

    # Iterate over tables in the document
    for table in tables:
        max_cols = max(len(row.cells) for row in table.rows)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                worksheet.write(index_row, j, cell.text.strip())
            # Fill the rest of the row with empty cells if necessary
            for j in range(len(row.cells), max_cols):
                worksheet.write(index_row, j, "")
            index_row += 1
        # Add a blank row after each table for separation
        index_row += 1

    # Close the workbook
    workbook.close()
    print(f'Tables extracted and saved to {excel_file_path}')

# Call the function to perform the conversion
docx_to_xlsx(docx_file_path, excel_file_path)