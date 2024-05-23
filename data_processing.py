import os
import pandas as pd
from pdf2docx import Converter
from docx import Document
import xlsxwriter

# Define base directory
base_dir = os.path.dirname(os.path.abspath(__file__))

# Paths to directories
pdf_dir = os.path.join(base_dir, 'mnt/data/pdfs')
docx_dir = os.path.join(base_dir, 'mnt/data/docx')
xlsx_dir = os.path.join(base_dir, 'mnt/data/xlsx')
csv_dir = os.path.join(base_dir, 'mnt/data/csv')

# Create directories if they do not exist
os.makedirs(docx_dir, exist_ok=True)
os.makedirs(xlsx_dir, exist_ok=True)
os.makedirs(csv_dir, exist_ok=True)

# Function to convert PDF to DOCX
def pdf_to_docx(pdf_path, docx_path):
    cv = Converter(pdf_path)
    cv.convert(docx_path, start=0, end=None)
    cv.close()

# Function to convert DOCX to XLSX
def docx_to_xlsx(docx_path, xlsx_path):
    document = Document(docx_path)
    tables = document.tables

    workbook = xlsxwriter.Workbook(xlsx_path)
    worksheet = workbook.add_worksheet("Combined Data")

    index_row = 0

    for table in tables:
        max_cols = max(len(row.cells) for row in table.rows)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                worksheet.write(index_row, j, cell.text.strip())
            for j in range(len(row.cells), max_cols):
                worksheet.write(index_row, j, "")
            index_row += 1
        index_row += 1

    workbook.close()

# Function to clean and save data from XLSX to CSV
def xlsx_to_csv(xlsx_path, base_name):
    df = pd.read_excel(xlsx_path, sheet_name='Combined Data', header=None)

    type1_rows = []
    type2_rows = []
    current_section = None

    def is_type1_heading(row):
        return row.str.contains('Date').any() and row.str.contains('Age').any() and row.str.contains('Life Event').any()

    def is_type2_heading(row):
        return row.str.contains('General Information').any()

    for index, row in df.iterrows():
        if is_type1_heading(row):
            current_section = 'type1'
            type1_rows.append(row.values.tolist())
        elif is_type2_heading(row):
            current_section = 'type2'
            type2_rows.append(['Heading', 'Value'])
        elif current_section == 'type1':
            type1_rows.append(row.values.tolist())
        elif current_section == 'type2':
            row_values = [v for v in row.values if pd.notna(v)]
            if len(row_values) == 2:
                type2_rows.append(row_values)
            elif len(row_values) > 2:
                type2_rows.append([row_values[0], ' '.join(map(str, row_values[1:]))])

    if type1_rows:
        type1_df = pd.DataFrame(type1_rows[1:], columns=type1_rows[0])
    else:
        type1_df = pd.DataFrame()

    if type2_rows:
        type2_df = pd.DataFrame(type2_rows[1:], columns=type2_rows[0])
    else:
        type2_df = pd.DataFrame()

    def clean_value(row):
        heading = row['Heading']
        value = row['Value']
        if heading in value:
            value = value.replace(heading, '').strip()
        return value

    if not type2_df.empty:
        type2_df['Value'] = type2_df.apply(clean_value, axis=1)

    type1_csv_path = os.path.join(csv_dir, f'Type1_{base_name}.csv')
    type2_csv_path = os.path.join(csv_dir, f'Type2_{base_name}.csv')

    type1_df.to_csv(type1_csv_path, index=False)
    type2_df.to_csv(type2_csv_path, index=False)

# Main script to process all PDFs
for pdf_filename in os.listdir(pdf_dir):
    if pdf_filename.endswith('.pdf'):
        base_name = os.path.splitext(pdf_filename)[0].replace(' ', '_').replace(',', '')
        pdf_path = os.path.join(pdf_dir, pdf_filename)
        docx_path = os.path.join(docx_dir, f'{base_name}.docx')
        xlsx_path = os.path.join(xlsx_dir, f'{base_name}.xlsx')

        # Convert PDF to DOCX
        pdf_to_docx(pdf_path, docx_path)

        # Convert DOCX to XLSX
        docx_to_xlsx(docx_path, xlsx_path)

        # Clean and save data from XLSX to CSV
        xlsx_to_csv(xlsx_path, base_name)

print("Processing completed.")